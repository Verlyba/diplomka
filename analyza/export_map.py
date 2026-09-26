"""Export polohy kostky na začátku + mapa 'kde se trefil a kde ne' (orchestrace × baseline)."""
import json, glob, os, sys, io, math
import numpy as np, openpyxl
from PIL import Image, ImageDraw, ImageFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
ROOT = "C:/Users/green/diplomka/"
SP = "C:/Users/green/AppData/Local/Temp/claude/C--Users-green-diplomka/3092d090-a58d-4566-873f-9ee439d802b4/scratchpad/"
XLSX = r"D:\_WinZakladniSlozky\zaznam_behu_2.xlsx"
OUT_XLSX = ROOT + "poznamky/poloha_kostky_start.xlsx"
OUT_PNG = ROOT + "poznamky/mapa_polohy_kostky.png"

def cube(p):
    a = np.array(Image.open(p).convert("RGB")).astype(int)
    m = (a[..., 1] > 80) & (a[..., 1] > a[..., 0] + 20) & (a[..., 1] > a[..., 2] + 20)
    ys, xs = np.nonzero(m)
    return None if len(xs) < 400 else (float(xs.mean()), float(ys.mean()))

# ── vstupy ──────────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook(XLSX)
wo, wbz = wb["Orchestrace"], wb["Baseline"]
o_user = {r: wo.cell(r, 13).value for r in range(5, 26)}
b_user = {r: wbz.cell(r, 8).value for r in range(5, 23)}
b_count = {r: wbz.cell(r, 7).value == "ANO" for r in range(5, 23)}
b_user_saved11 = b_user[11]
b_user[11] = "NE"                                  # oprava uživatele (v uloženém souboru zatím ANO)

demos = np.array(json.load(open(SP + "demos_pose_pos.json"))["P"])
base_est = json.load(open(SP + "base_pos.json"))

runs_all = sorted(os.path.basename(f)[:-5] for f in glob.glob(ROOT + "runs/20260926-*.json") if os.path.basename(f)[:-5] >= "20260926-164147")
info = {}
for rid in runs_all:
    r = json.load(open(ROOT + f"runs/{rid}.json", encoding="utf-8"))
    init = (r.get("initial_images") or [None])[0]
    c = cube(ROOT + init) if init else None
    cs = "cs100" if "diplomka_1" in json.dumps(r["catalog"]) else "cs15"
    catches = [s for s in r["steps"] if s["step"] == "catch_cube"]
    info[rid] = dict(pos=c, cs=cs, sys="ANO" if r["success"] else "NE", dur=r["duration_s"], replans=max([s.get("replan", 0) for s in r["steps"]] or [0]),
                     n_catch=len(catches), first=bool(r["steps"]) and r["steps"][0]["step"] == "catch_cube" and r["steps"][0]["success"])
rec = runs_all[:21]                                   # ty, které jsou v sešitu (řádky 5–25)
extra = runs_all[21:]                                 # novější běhy (po odebrání re-snapshotu), v sešitu zatím nejsou
orch15 = rec[:16]; orch100 = rec[16:]
row_of = {rid: 5 + i for i, rid in enumerate(rec)}

def ncov(p, rad=60): return int((np.hypot(demos[:, 0] - p[0], demos[:, 1] - p[1]) < rad).sum())

# baseline: započítané pokusy v pořadí -> pozice 1..16 (stejné pořadí jako orchestrace, viz test v konverzaci)
b_rows = [r for r in range(5, 23) if b_count[r]]
assert len(b_rows) == 16
UNSURE = {5, 6, 7, 8, 9}                                # pozice 5–9 (bloky u spodního okraje) – přiřazení uvnitř bloku nejisté
# cs100 opakuje pozice řádků 5, 6, 7, 14, 8 -> pozice 1, 2, 3, 10, 4
cs100_pos = {}
for rid in orch100:
    p = info[rid]["pos"]; d = [math.hypot(p[0] - info[o]["pos"][0], p[1] - info[o]["pos"][1]) for o in orch15]
    cs100_pos[rid] = int(np.argmin(d)) + 1

rows = []
for k in range(16):
    rid = orch15[k]; i = info[rid]; b = b_rows[k]; e = base_est[str(b)]
    c100 = [x for x in orch100 if cs100_pos[x] == k + 1]
    rows.append(dict(k=k + 1, x=i["pos"][0], y=i["pos"][1], cov=ncov(i["pos"]), o_row=row_of[rid], o_id=rid, o_user=o_user[row_of[rid]], o_sys=i["sys"],
                     o_att=i["n_catch"], o_first=i["first"], o_rep=i["replans"], o_dur=i["dur"],
                     b_row=b, b_user=b_user[b], b_file=f"20260926-{['184956','185206','185318','185530','185704','185846','190024','190215','190328','190437','190546','190737','190910','191013','191132','191257','191504','191623'][b-5]}",
                     b_est=(e["pos"] if e else None), b_uncertain=(k + 1) in UNSURE,
                     c_id=(c100[0] if c100 else ""), c_row=(row_of[c100[0]] if c100 else ""), c_user=(o_user[row_of[c100[0]]] if c100 else "")))

# ── statistiky ──────────────────────────────────────────────────────────────
A = [r["o_user"] == "ANO" for r in rows]; B = [r["b_user"] == "ANO" for r in rows]
both = sum(a and b for a, b in zip(A, B)); oo = sum(a and not b for a, b in zip(A, B)); bo = sum(b and not a for a, b in zip(A, B)); nn = sum(not a and not b for a, b in zip(A, B))
def mcnemar(b, c):
    n = b + c; k = min(b, c); p = sum(math.comb(n, i) for i in range(0, k + 1)) / 2 ** n * 2
    return min(1.0, p)
print(f"orchestrace cs15 {sum(A)}/16 | baseline {sum(B)}/16 | obě {both}, jen orchestrace {oo}, jen baseline {bo}, ani jedna {nn} | McNemar p={mcnemar(oo, bo):.2f}")
blk = [r for r in rows if r["b_uncertain"]]
print("blok pozic 5–9 (pořadí baseline nejisté): orchestrace", sum(r["o_user"] == "ANO" for r in blk), "/5, baseline", sum(r["b_user"] == "ANO" for r in blk), "/5")
print("mimo blok (11 poloh): orchestrace", sum(r["o_user"] == "ANO" for r in rows if not r["b_uncertain"]), "/11, baseline", sum(r["b_user"] == "ANO" for r in rows if not r["b_uncertain"]), "/11")
print("první pokus orchestrace úspěšný a fakticky správný:", sum(r["o_first"] and r["o_user"] == "ANO" and r["o_rep"] == 0 for r in rows), "/16")

# ── obrázek ─────────────────────────────────────────────────────────────────
S = 2
imgs = [np.array(Image.open(ROOT + json.load(open(ROOT + f"runs/{r}.json", encoding="utf-8"))["initial_images"][0]).convert("RGB")) for r in rec]
bg = np.median(np.stack(imgs), axis=0).astype(np.uint8)          # kostka se mezi snímky mění -> medián ji odstraní
bgim = Image.fromarray(bg).resize((640 * S, 480 * S), Image.LANCZOS)
bgim = Image.blend(bgim, Image.new("RGB", bgim.size, "white"), 0.45)
W, H = bgim.size; PAD_B = 250
canvas = Image.new("RGB", (W, H + PAD_B), "white"); canvas.paste(bgim, (0, 0))
d = ImageDraw.Draw(canvas, "RGBA")
def font(sz, bold=False):
    for f in (["C:/Windows/Fonts/arialbd.ttf"] if bold else []) + ["C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/segoeui.ttf"]:
        try: return ImageFont.truetype(f, sz)
        except Exception: pass
    return ImageFont.load_default()
F, FB, FS = font(22), font(24, True), font(18)
GREEN, RED, GRAY = (46, 158, 79, 255), (214, 69, 69, 255), (120, 120, 120, 255)
for x, y in demos: d.ellipse([x * S - 3, y * S - 3, x * S + 3, y * S + 3], fill=(90, 90, 90, 110))
R = 25
# nejisté bloky
for r in rows:
    if r["b_uncertain"]:
        cx, cy = r["x"] * S, r["y"] * S
        d.ellipse([cx - R - 9, cy - R - 9, cx + R + 9, cy + R + 9], outline=(60, 60, 60, 255), width=3)
for r in rows:
    cx, cy = r["x"] * S, r["y"] * S
    col_o = GREEN if r["o_user"] == "ANO" else RED; col_b = GREEN if r["b_user"] == "ANO" else RED
    d.pieslice([cx - R, cy - R, cx + R, cy + R], 90, 270, fill=col_o)         # levá polovina = orchestrace
    d.pieslice([cx - R, cy - R, cx + R, cy + R], 270, 90, fill=col_b)         # pravá polovina = baseline
    d.ellipse([cx - R, cy - R, cx + R, cy + R], outline=(255, 255, 255, 255), width=3)
    lab = str(r["k"]); tw = d.textlength(lab, font=FB)
    d.rectangle([cx - tw / 2 - 4, cy - R - 30, cx + tw / 2 + 4, cy - R - 4], fill=(255, 255, 255, 230))
    d.text((cx - tw / 2, cy - R - 31), lab, fill=(20, 20, 20, 255), font=FB)
    att = f"{r['o_att']}×" if r["o_user"] == "ANO" else f"{r['o_att']}× ne"
    tw = d.textlength(att, font=FS)
    d.text((cx - tw / 2, cy + R + 4), att, fill=(20, 20, 20, 255), font=FS)
# cs100 – malý kosočtverec vpravo od značky
for rid in orch100:
    k = cs100_pos[rid] - 1; r = rows[k]; cx, cy = r["x"] * S + R + 20, r["y"] * S
    col = GREEN if o_user[row_of[rid]] == "ANO" else RED
    d.polygon([(cx, cy - 13), (cx + 13, cy), (cx, cy + 13), (cx - 13, cy)], fill=col, outline=(255, 255, 255, 255))
# legenda
y0 = H + 12
d.text((16, y0), "Poloha kostky na začátku (ze snímku, 640×480 px) a výsledek podle úsudku uživatele", fill=(20, 20, 20, 255), font=FB)
def lg(x, y, half_l, half_r, text):
    d.pieslice([x, y, x + 36, y + 36], 90, 270, fill=half_l); d.pieslice([x, y, x + 36, y + 36], 270, 90, fill=half_r)
    d.ellipse([x, y, x + 36, y + 36], outline=(120, 120, 120, 255), width=2); d.text((x + 50, y + 5), text, fill=(20, 20, 20, 255), font=F)
lg(16, y0 + 44, GREEN, RED, "levá půlka = orchestrace cs15 (s re-plány), pravá půlka = baseline (jeden pokus)")
d.text((66, y0 + 86), "zelená = trefil (úsudek uživatele), červená = netrefil", fill=(20, 20, 20, 255), font=F)
d.polygon([(28, y0 + 135), (41, y0 + 122), (54, y0 + 135), (41, y0 + 148)], fill=GREEN)
d.text((66, y0 + 124), "kosočtverec = orchestrace cs100 (zkušební běhy, jen 5 poloh)", fill=(20, 20, 20, 255), font=F)
d.text((66, y0 + 160), "číslo pod značkou = počet úchopů (pokusů catch_cube) orchestrace, „ne“ = ani po nich netrefila", fill=(20, 20, 20, 255), font=F)
d.ellipse([16, y0 + 196, 40, y0 + 220], outline=(60, 60, 60, 255), width=3)
d.text((66, y0 + 196), "kroužek = přiřazení baseline uvnitř bloku poloh 5–9 je nejisté", fill=(20, 20, 20, 255), font=F)
d.ellipse([900, y0 + 205, 906, y0 + 211], fill=(90, 90, 90, 200)); d.text((914, y0 + 196), "demonstrace (120)", fill=(20, 20, 20, 255), font=F)
canvas.save(OUT_PNG)
print("obrázek:", OUT_PNG, canvas.size)

# ── xlsx ────────────────────────────────────────────────────────────────────
Hf = PatternFill("solid", fgColor="1F4E78"); Hn = Font(name="Arial", size=10, bold=True, color="FFFFFF"); Nn = Font(name="Arial", size=10)
thin = Side(style="thin", color="D9D9D9"); BR = Border(left=thin, right=thin, top=thin, bottom=thin)
okf, nef = PatternFill("solid", fgColor="E2EFDA"), PatternFill("solid", fgColor="FCE4D6")
wx = openpyxl.Workbook(); ws = wx.active; ws.title = "Polohy"
ws["A1"] = "Poloha kostky na začátku a výsledek (16 poloh, cs15)"; ws["A1"].font = Font(name="Arial", size=14, bold=True)
ws["A2"] = ("Poloha = střed zelené kostky na úvodním snímku orchestrace (px v 640×480). Baseline žádný snímek neukládá: jeho pokusy jsou přiřazeny ke stejným polohám podle "
            "pořadí (potvrzeno odhadem polohy z pozice ramene, viz list Baseline_odhad); v bloku poloh 5–9 je přiřazení uvnitř bloku nejisté. Úsudek = váš (řádek 11 baseline podle opravy: NE).")
ws["A2"].font = Font(name="Arial", size=9, italic=True, color="595959"); ws["A2"].alignment = Alignment(wrap_text=True); ws.merge_cells("A2:R2"); ws.row_dimensions[2].height = 42
hdr = ["Poloha č.", "x [px]", "y [px]", "Demonstrací do 60 px", "Orch. řádek", "Orch. ID běhu", "Orch. úsudek", "Orch. úchopů", "Orch. re-plánů", "Orch. doba [s]",
       "Baseline řádek", "Baseline soubor", "Baseline úsudek", "Baseline odhad polohy z ramene [px]", "Přiřazení baseline", "cs100 řádek", "cs100 ID", "cs100 úsudek"]
for i, h in enumerate(hdr, 1):
    c = ws.cell(4, i, h); c.font = Hn; c.fill = Hf; c.border = BR; c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
ws.row_dimensions[4].height = 44
for j, r in enumerate(rows):
    vals = [r["k"], round(r["x"]), round(r["y"]), r["cov"], r["o_row"], r["o_id"], r["o_user"], r["o_att"], r["o_rep"], r["o_dur"], r["b_row"], r["b_file"], r["b_user"],
            (f"({r['b_est'][0]:.0f}, {r['b_est'][1]:.0f})" if r["b_est"] else "-"), "nejisté (blok 5–9)" if r["b_uncertain"] else "jisté", r["c_row"], r["c_id"], r["c_user"]]
    for i, v in enumerate(vals, 1):
        c = ws.cell(5 + j, i, v); c.font = Nn; c.border = BR; c.alignment = Alignment(vertical="center", horizontal="center")
    ws.cell(5 + j, 7).fill = okf if r["o_user"] == "ANO" else nef; ws.cell(5 + j, 13).fill = okf if r["b_user"] == "ANO" else nef
    if r["c_user"]: ws.cell(5 + j, 18).fill = okf if r["c_user"] == "ANO" else nef
n0 = 5 + len(rows) + 1
ws.cell(n0, 1, "Souhrn").font = Font(name="Arial", size=11, bold=True)
sm = [("Orchestrace cs15 úspěšných", f"{sum(A)}/16"), ("Baseline úspěšných", f"{sum(B)}/16"), ("Obě trefily", both), ("Jen orchestrace", oo), ("Jen baseline", bo), ("Ani jedna", nn),
      ("McNemar (dvoustranný, přesný)", round(mcnemar(oo, bo), 2)), ("Blok poloh 5–9: orchestrace / baseline", f"{sum(r['o_user']=='ANO' for r in blk)}/5 , {sum(r['b_user']=='ANO' for r in blk)}/5"),
      ("Mimo blok (11 poloh): orchestrace / baseline", f"{sum(r['o_user']=='ANO' for r in rows if not r['b_uncertain'])}/11 , {sum(r['b_user']=='ANO' for r in rows if not r['b_uncertain'])}/11")]
for j, (a, b) in enumerate(sm, 1):
    ws.cell(n0 + j, 1, a).font = Nn; ws.cell(n0 + j, 4, b).font = Nn
for i, w in enumerate([9, 8, 8, 12, 9, 18, 10, 10, 10, 10, 10, 20, 10, 20, 18, 9, 18, 10], 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A5"

w2 = wx.create_sheet("Všechny běhy")
h2 = ["Zdroj", "Řádek v sešitu", "ID", "Sada", "x [px]", "y [px]", "Úsudek uživatele", "Verdikt systému", "Úchopů", "Re-plánů", "Doba [s]", "Poznámka"]
for i, h in enumerate(h2, 1):
    c = w2.cell(1, i, h); c.font = Hn; c.fill = Hf; c.border = BR; c.alignment = Alignment(wrap_text=True, horizontal="center")
r2 = 2
for rid in runs_all:
    i = info[rid]; rw = row_of.get(rid, "")
    note = "" if rid in row_of else "novější běh (po odebrání re-snapshotu), v sešitu zatím není"
    vals = ["orchestrace", rw, rid, i["cs"], round(i["pos"][0]) if i["pos"] else "", round(i["pos"][1]) if i["pos"] else "", o_user.get(rw, ""), i["sys"], i["n_catch"], i["replans"], i["dur"], note]
    for j, v in enumerate(vals, 1):
        c = w2.cell(r2, j, v); c.font = Nn; c.border = BR
    r2 += 1
w2.freeze_panes = "A2"
for i, w in enumerate([12, 10, 20, 8, 8, 8, 12, 12, 8, 9, 9, 52], 1): w2.column_dimensions[get_column_letter(i)].width = w

w3 = wx.create_sheet("Baseline_odhad")
w3["A1"] = ("Odhad polohy kostky u baseline z pozice ramene v okamžiku, kdy začne zavírat gripper (kNN nad 120 demonstracemi). Ověření na 18 skutečných úspěšných úchopech orchestrace: "
            "medián chyby 22 px, max 82 px. U netrefených pokusů míří rameno vedle kostky, takže odhad je horší (spodní okraj desky jde systematicky krátký).")
w3["A1"].alignment = Alignment(wrap_text=True); w3["A1"].font = Font(name="Arial", size=9, italic=True); w3.merge_cells("A1:H1"); w3.row_dimensions[1].height = 44
h3 = ["Baseline řádek", "Soubor", "Počítá se", "Úsudek", "Zavírá gripper v [s]", "Odhad x [px]", "Odhad y [px]", "Vzdálenost k nejbližším demonstracím v pozici ramene (menší = věrohodnější)"]
for i, h in enumerate(h3, 1):
    c = w3.cell(3, i, h); c.font = Hn; c.fill = Hf; c.border = BR; c.alignment = Alignment(wrap_text=True, horizontal="center")
w3.row_dimensions[3].height = 56
files = ['184956','185206','185318','185530','185704','185846','190024','190215','190328','190437','190546','190737','190910','191013','191132','191257','191504','191623']
for j, r in enumerate(range(5, 23)):
    e = base_est[str(r)]
    vals = [r, f"20260926-{files[j]}", "ano" if b_count[r] else "ne", b_user[r], round(e["t_close"], 1) if e else "-", round(e["pos"][0]) if e else "-", round(e["pos"][1]) if e else "-", round(e["dist"], 1) if e else "-"]
    for i, v in enumerate(vals, 1):
        c = w3.cell(4 + j, i, v); c.font = Nn; c.border = BR
for i, w in enumerate([10, 20, 10, 10, 14, 12, 12, 44], 1): w3.column_dimensions[get_column_letter(i)].width = w
wx.save(OUT_XLSX)
print("xlsx:", OUT_XLSX, "| (uložená verze řádku 11 baseline v tvém sešitu:", b_user_saved11, ")")
json.dump([{k: v for k, v in r.items() if k != "b_est"} | {"b_est": r["b_est"]} for r in rows], open(SP + "rows.json", "w"), ensure_ascii=False, default=str)
