"""Doplní zaznam_behu_2.xlsx z runs/*.json a telemetry/*.jsonl a přidá list Kontrola.
Ručně vyplněné sloupce (POCITA SE?, VAS USUDEK, Poznamka, Model(y)) se NEMĚNÍ."""
import json, glob, os, sys, io, math, time, copy
import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
sys.path.insert(0, "C:/Users/green/diplomka")
from release_detect import ReleaseTracker

ROOT = "C:/Users/green/diplomka/"
SP = "C:/Users/green/AppData/Local/Temp/claude/C--Users-green-diplomka/3092d090-a58d-4566-873f-9ee439d802b4/scratchpad/"
SRC = r"D:\_WinZakladniSlozky\zaznam_behu_2.xlsx"
DST = r"D:\_WinZakladniSlozky\zaznam_behu_2_doplneno.xlsx"

# ── mapování ────────────────────────────────────────────────────────────────
ORCH_RUNS = sorted(os.path.basename(f)[:-5] for f in glob.glob(ROOT + "runs/20260926-*.json")
                   if os.path.basename(f)[:-5] >= "20260926-164147")
BASE_FILES = ["184956", "185206", "185318", "185530", "185704", "185846", "190024", "190215", "190328",
              "190437", "190546", "190737", "190910", "191013", "191132", "191257", "191504", "191623"]
assert len(ORCH_RUNS) == 21 and len(BASE_FILES) == 18

def stamp(name):  # "20260926-164147" -> ("2026-09-26", "16:41")
    d, t = name[:8], name[9:15]
    return f"{d[:4]}-{d[4:6]}-{d[6:]}", f"{t[:2]}:{t[2:4]}"

def ep_of(pp):
    import re
    m = re.search(r"_(\d+)ep_", (pp or "").replace("\\", "/"))
    return m.group(1) if m else ""

def load_jsonl(f):
    out = []
    for l in open(f, encoding="utf-8", errors="replace"):
        l = l.strip()
        if l:
            try: out.append(json.loads(l))
            except Exception: pass
    return out

wb = openpyxl.load_workbook(SRC)
orig = openpyxl.load_workbook(SRC)          # pro pozdější kontrolu, že se ruční sloupce nezměnily

# ── Orchestrace ────────────────────────────────────────────────────────────
wo = wb["Orchestrace"]
cube_pos = json.load(open(SP + "cube_pos.json"))
cov = json.load(open(SP + "coverage.json"))
orch_info = {}
for i, rid in enumerate(ORCH_RUNS):
    row = 5 + i
    r = json.load(open(ROOT + f"runs/{rid}.json", encoding="utf-8"))
    cat = r.get("catalog") or []
    eps = sorted({ep_of(s.get("policy_path")) for s in cat if ep_of(s.get("policy_path"))})
    assert eps == ["120"], (rid, eps)
    cs = r.get("cost_summary") or {}
    st = r["steps"]
    cond = "Orchestrace (bez re-planu)" if r["config"].get("max_replans") == 0 else "Orchestrace"
    d, t = stamp(rid)
    vals = {1: rid, 2: d, 3: t, 4: cond, 6: "120", 7: round(float(r["duration_s"]), 1),
            8: max((s.get("replan") or 0) for s in st) if st else 0,
            9: cs.get("planner_calls"), 10: cs.get("inspector_calls"), 11: "ANO" if r.get("success") else "NE"}
    for col, v in vals.items():
        assert wo.cell(row, col).value in (None, ""), (row, col, wo.cell(row, col).value)
        wo.cell(row, col, v)
    seq = " ".join(f"{ {'catch_cube':'C','carry_cube':'K','homing':'H'}.get(s['step'], s['step'][:1].upper())}{'+' if s['success'] else '-'}" for s in st)
    orch_info[row] = dict(id=rid, seq=seq, sys=vals[11], dur=vals[7], replans=vals[8], ceo=vals[9], vlm=vals[10],
                          model="cs100 (diplomka_1)" if "diplomka_1" in json.dumps(cat) else "cs15 (diplomka_2)",
                          r=r)

# ── Baseline ───────────────────────────────────────────────────────────────
wbz = wb["Baseline"]
LIMIT = 300.0
base_info = {}
for i, fn in enumerate(BASE_FILES):
    row = 5 + i
    stem = f"20260926-{fn}"
    ev = load_jsonl(ROOT + f"telemetry/{stem}.jsonl")
    ds = next(e for e in ev if e["event"] == "daemon_start")
    pol = os.path.basename(ds["policy_path"].replace("\\", "/"))
    sts = [e for e in ev if e["event"] == "task_started"]
    assert len(sts) == 1, (stem, len(sts))
    dur = round(ev[-1]["t"] - sts[0]["t"], 1)                     # stejně jako build_runs_log.py: start -> konec souboru
    d, t = stamp(stem)
    vals = {1: f"{stem}#1", 2: d, 3: t, 4: pol, 5: ep_of(pol), 6: dur}
    for col, v in vals.items():
        assert wbz.cell(row, col).value in (None, ""), (row, col)
        wbz.cell(row, col, v)
    # indicie z telemetrie (jen do listu Kontrola)
    tk = [e for e in ev if e["event"] == "tick" and e.get("state") == "RUNNING" and e.get("joints")]
    seg = [e for e in tk if e["t"] >= sts[0]["t"]]
    j0 = seg[0]["joints"]
    move = max(max(abs(a - b) for a, b in zip(e["joints"][:5], j0[:5])) for e in seg)
    tr = ReleaseTracker(0.66 * LIMIT, 0.5 * LIMIT, 1, 2, 3.0)
    g = rel = None; prev = None; held = 0
    for e in seg:
        step = abs(e["joints"][5] - prev) if prev is not None else float("inf"); prev = e["joints"][5]
        if e["t"] - sts[0]["t"] < 0.75: continue
        wg, wr = tr.gripped, tr.released
        tr.update(e["load"] - (e.get("baseline") or 0.0), step, abs(e.get("slope") or 0) < 30)
        if tr.holding: held += 1
        if tr.gripped and not wg and g is None: g = e["t"] - sts[0]["t"]
        if tr.released and not wr and rel is None: rel = e["t"] - sts[0]["t"]
    base_info[row] = dict(file=stem, dur=dur, move=move, grasp=g, held=held * 0.2, release=rel,
                          maxload=max(e["load"] for e in seg))

# ── Souhrn (odpovídá vzorcům v listu Souhrn) ────────────────────────────────
def wilson(k, n, z=1.959963985):
    p = k / n; c = (p + z * z / (2 * n)) / (1 + z * z / n)
    m = (z / (1 + z * z / n)) * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)); return max(0, c - m), min(1, c + m)
summ = {}
def count(ws, cond_col, cond, ep_col, pocita_col, usud_col, dur_col):
    n = k = 0; durs = []
    for r in range(5, ws.max_row + 1):
        if ws.cell(r, ep_col).value != "120": continue
        if cond_col and ws.cell(r, cond_col).value != cond: continue
        if ws.cell(r, pocita_col).value == "ANO":
            n += 1; k += ws.cell(r, usud_col).value == "ANO"; durs.append(ws.cell(r, dur_col).value)
    return n, k, (sum(durs) / len(durs) if durs else None)
summ["Baseline 120"] = count(wbz, None, None, 5, 7, 8, 6)
summ["Orchestrace 120"] = count(wo, 4, "Orchestrace", 6, 12, 13, 7)

# ── list Kontrola ───────────────────────────────────────────────────────────
FONT = "Arial"
H_FILL = PatternFill("solid", fgColor="1F4E78"); H_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
N_FONT = Font(name=FONT, size=10); B_FONT = Font(name=FONT, size=10, bold=True)
T_FONT = Font(name=FONT, size=14, bold=True); S_FONT = Font(name=FONT, size=12, bold=True, color="1F4E78")
NOTE = Font(name=FONT, size=9, italic=True, color="595959")
WARN = PatternFill("solid", fgColor="FCE4D6"); OKF = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="D9D9D9"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
wk = wb.create_sheet("Kontrola")
wk.sheet_view.showGridLines = False
wk["A1"] = "Kontrola doplněných dat"; wk["A1"].font = T_FONT
wk["A2"] = ("Vygenerováno z runs/*.json, telemetry/*.jsonl a uložených snímků. Vaše sloupce (POCITA SE?, VAS USUDEK, Poznamka, Model(y)) "
            "zůstaly beze změny. 'Závěrečný snímek' je MOJE čtení uložených snímků (horní kamera + zápěstí), ne měření.")
wk["A2"].font = NOTE; wk["A2"].alignment = Alignment(wrap_text=True); wk.merge_cells("A2:N2"); wk.row_dimensions[2].height = 30

def header(r, cols):
    for i, h in enumerate(cols, 1):
        c = wk.cell(r, i, h); c.font = H_FONT; c.fill = H_FILL; c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

def put(r, vals, fill=None, bold_cols=()):
    for i, v in enumerate(vals, 1):
        c = wk.cell(r, i, v); c.font = B_FONT if i in bold_cols else N_FONT; c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if fill: c.fill = fill

# --- 1) nálezy
r0 = 4
wk.cell(r0, 1, "1) Hlavní nálezy").font = S_FONT
FIND = [
 ("Mapování orchestrace",
  "Řádky 5–25 = běhy 20260926-164147 … 20260926-181830 (21 běhů: 16× cs15 = diplomka_2, 5× cs100 = diplomka_1). Potvrzeno dvěma nezávislými způsoby: "
  "(a) poznámky v řádcích 5–8 sedí na běhy 164147, 164738, 165409, 170037 (třetí úchop + chybný homing + done_check; opakovaný homing na konci; úspěch až po pátém re-plánu); "
  "(b) polohy kostky u cs100 řádků 21, 22, 23, 24, 25 opakují polohy řádků 5, 6, 7, 14, 8 s odchylkou do 8 px. 16 poloh cs15 je všech 16 navzájem různých."),
 ("'4. běh dne'",
  "V souborech runs/ je 4. běh dne 20260926-163206 (ten 'nesmyslný', kde homing chybně selhal) a ten v tabulce NENÍ. Tabulka začíná 164147, což je 5. soubor dne "
  "(1. je ranní 105658). Pokud jsi ranní běh nepočítal, je 164147 tvůj 4. běh a vše sedí."),
 ("Mapování baseline",
  "Řádky 5–22 = telemetrie 20260926-184956 … 191623 (18 pokusů, všechny diplomka_2_120ep_act). Nezahrnuto: 184825 (3 kratší pokusy před řádkem 5), 191238 (démon bez úlohy), "
  "215517 z 25. 9. Kotva mapování: řádek 8 je jediný pokus, kde se rameno nehnulo (0,4° za 10,6 s, zátěž 20) = 'přetížený motor'. Ano, uložil se (20260926-185530)."),
 ("⚠ Baseline řádek 11",
  "Úsudek ANO odporuje poznámce ('špatně najel … jen odsouval, ukončeno mnou') i telemetrii (žádné trvalé sevření, 40 s do ručního ukončení). Pravděpodobně má být NE. "
  "Nezměnil jsem to, je to Váš úsudek. Se změnou na NE vychází baseline 7/16 místo 8/16."),
 ("Úsudky vs. snímky",
  "U všech 21 řádků orchestrace sedí Váš úsudek se závěrečným snímkem (kostka v misce u ANO, mimo misku u NE), podle mého čtení snímků."),
 ("Verdikt systému vs. Váš úsudek",
  "Rozdíl u 4 běhů z 21: řádek 7 (systém NE, Vy ANO – kostka je v misce, běh skončil chybou 'homing podruhé za sebou') a řádky 16, 17, 25 (systém ANO, Vy NE – kostka není v misce). "
  "Mezi 16 započítanými (cs15) sedí 13/16. Systém tedy nadhodnocuje (12 ANO vs Vašich 11 ANO u cs15) a to i o 'špatné důvody' – viz další řádky."),
 ("Řádky 16, 17",
  "Úchop: fyzika DENY (zátěž nepřešla práh), přesto inspektor napsal 'drží kostku' → SUCCESS. Přenos: fyzika NONE (nic se nesevřelo), inspektor 'gripper je v misce' → SUCCESS. "
  "Na snímcích leží kostka na desce mimo misku. Chyba je čistě v inspektorovi, fyzika ji tu dvakrát správně nepotvrdila."),
 ("Řádek 25 (cs100)",
  "Přenos vyhodnotila jako úspěch FYZIKA (uvolnění po dojezdu: 31 tiků klidu, gripper otevřen na 40,7°), inspektor byl UNCLEAR. Kostka ale spadla mimo misku "
  "(leží na spodním okraji desky). Nové čtení 'uvolnění' říká jen 'čelisti pustily', ne 'kam'. Do měření cs15 to nezasáhlo (řádek nemá POCITA SE), ale je to slabina pravidla."),
 ("Homing chybně selhává",
  "Homing FAIL při kostce v misce a rameni v domovské poloze: běhy 164147 (#7), 165409 (#9, #10), i nezapočítaný 163206 (#3). Inspektor píše 'gripper drží předmět'. "
  "Fyzika u nich CONFIRM, ale při konfliktu CONFIRM + jistý FAIL rozhoduje inspektor → krok selže → další plán, zbytečné kroky navíc a delší běh. Ovlivňuje Dobu běhu a Re-plánů, ne Váš úsudek."),
 ("Pokrytí demonstracemi",
  "Všech 21 poloh má demonstrace poblíž: nejbližší 3–22 px (kostka má ~35 px), 3–13 demonstrací do 40 px. Nejřidší je řádek 12 (4 demo do 60 px) a řádek 20 (5). "
  "Poznámky 'málo trénovacích dat' u řádků 12 a 14: u 12 se potvrzuje (nejřidší místo), u 14 ne (16 demo do 60 px, jako běžné místo). "
  "Řádek 20 (řídké místo) uspěl na první pokus, řádek 12 (stejně řídké) neuspěl ani jednou – řídké místo tedy neznamená jistý neúspěch. Nejhůř pokrytých 8 poloh: 3/8 úspěšných prvních úchopů (podle systému), nejlépe pokrytých 8: 4/8."),
 ("Doba pokusu u baseline",
  "Je to čas od SET_TASK po poslední záznam v telemetrii, tj. do chvíle, kdy jsi daemon zastavil. U úspěchů zahrnuje stání po dokončení. Přesnější odhad je sloupec 'Puštění v [s]' níže "
  "(kdy se gripper naposledy pustil = kostka v misce podle zátěže; mediánově ~10 s). Do Souhrnu ale jde původní Doba pokusu."),
 ("cs100 řádky",
  "Všech 5 má POCITA SE = NE, takže Souhrn počítá jen cs15. Baseline běžel také na cs15 (diplomka_2_120ep_act), takže srovnání je cs15 vs cs15."),
]
r = r0 + 1
header(r, ["Téma", "Nález"] + [""] * 12); wk.merge_cells(start_row=r, start_column=2, end_row=r, end_column=14)
for t, txt in FIND:
    r += 1
    put(r, [t, txt] + [""] * 12, fill=WARN if t.startswith("⚠") else None, bold_cols=(1,))
    wk.merge_cells(start_row=r, start_column=2, end_row=r, end_column=14)
    wk.row_dimensions[r].height = max(30, 15 * math.ceil(len(txt) / 150))

# --- 2) orchestrace po řádcích
FINAL = {5:"v misce",6:"v misce",7:"v misce",8:"v misce",9:"v misce",10:"v misce",11:"mimo misku (spodní okraj desky)",12:"mimo misku (levý dolní roh)",
         13:"v misce",14:"mimo misku (u horního okraje misky)",15:"v misce",16:"mimo misku (na desce)",17:"mimo misku (na desce)",18:"v misce",19:"v misce",
         20:"v misce",21:"v misce",22:"v misce",23:"v misce",24:"v misce",25:"mimo misku (spodní okraj desky)"}
COMM = {
 5:"Sedí: 3. úchop (C- C- C+), poslední homing chybně FAIL ('gripper drží'), done_check přesto potvrdil cíl.",
 6:"Sedí: 3. úchop, poslední homing OK. Posun kostky mezi pokusy jsem numericky neověřoval.",
 7:"Sedí: 3. úchop; na konci 2× homing FAIL (chybně) → plánovač skončil chybou. Systém=NE, kostka je v misce → Váš ANO správně.",
 8:"Sedí, ale 5 neúspěšných úchopů + úspěch na 6. pokusu (= 5. re-plán, poznámka říká 'pátý pokus').",
 9:"Sedí: 1. pokus (C+ K+ H+).", 10:"Sedí: 1. pokus (C+ K+ H+).",
 11:"Sedí: 6 neúspěšných úchopů. Kostka ležela nakonec na spodním okraji desky.",
 12:"Sedí: všechny úchopy neúspěšné. Nejřidší pokrytí demonstracemi (4 do 60 px) – 'málo dat' se potvrzuje.",
 13:"Sedí: úspěch na 3. úchop (C- C- C+).",
 14:"Sedí: všechny úchopy neúspěšné. Pokrytí je ale běžné (16 demo do 60 px) → 'málo dat' se nepotvrzuje; místo je nejdál od ramene u misky.",
 15:"Bez poznámky. Log: C- H+ C+ K+ H+ (2. úchop).",
 16:"Sedí: inspektor 2× chybně (úchop i přenos) → systém ANO, ve skutečnosti kostka na desce.",
 17:"Sedí: totéž jako řádek 16 (C+ K- K+ H+ a kostka mimo misku).",
 18:"Sedí: 1. pokus.", 19:"Bez poznámky. Log: C- H+ C+ K+ H+ (2. úchop).",
 20:"Sedí: 1. pokus. Řídká oblast (3 demo do 40 px); protější řídký roh (řádek 12) neuspěl.",
 21:"cs100. Sedí: 1. pokus, bez homingu (prompt ho ještě neobsahoval). Stejné místo jako řádek 5.",
 22:"cs100. Sedí: 1. pokus, bez homingu. Stejné místo jako řádek 6.",
 23:"cs100. Sedí: 1. úchop neúspěšný, 2. úspěšný. Poslední přenos K- (systém), přitom je kostka v misce. Stejné místo jako řádek 7.",
 24:"cs100. Sedí: 1. pokus, stejné místo jako řádek 14 (kde cs15 neuspěl ani jednou).",
 25:"cs100. Sedí (NE), ale systém=ANO je falešný: přenos potvrdila fyzika, kostka je ovšem mimo misku. Stejné místo jako řádek 8.",
}
r += 2
wk.cell(r, 1, "2) Orchestrace – řádek po řádku").font = S_FONT
r += 1
header(r, ["Řádek", "ID běhu", "Sada", "Kroky (C=catch K=carry H=homing)", "Doba [s]", "Re-plány", "Poloha kostky na startu [px]",
           "Nejbližší demo [px]", "Demo do 60 px", "Systém", "Váš úsudek", "Závěrečný snímek (moje čtení)", "Sedí úsudek?", "Kontrola poznámky"])
wk.row_dimensions[r].height = 42
for row in range(5, 26):
    o = orch_info[row]
    user = wo.cell(row, 13).value
    ok = (user == "ANO") == (FINAL[row] == "v misce")
    c = cube_pos[o["id"]]; cv = cov[str(row)]
    r += 1
    put(r, [row, o["id"], o["model"], o["seq"], o["dur"], o["replans"], f"({c[0]:.0f}, {c[1]:.0f})", round(cv["nearest"]), cv["n60"],
            o["sys"], user, FINAL[row], "ano" if ok else "NE – ZKONTROLUJ", COMM[row]],
        fill=None if ok else WARN)
    if (o["sys"] == "ANO") != (user == "ANO"):
        wk.cell(r, 10).fill = WARN
    wk.row_dimensions[r].height = 45

# --- 3) baseline po řádcích
BCOMM = {
 5:"Sedí: sevření a puštění.", 6:"Krátké sevření (1,2 s) – s poznámkou 'netrefil' se nepere (kontakt s kostkou, ne přenos).",
 7:"Krátké sevření (1,6 s) – sedí s 'odstřelil stranou'.", 8:"Sedí: rameno se nehnulo (0,4°), zátěž 20 – přetížený motor. Soubor se uložil.",
 9:"Sedí: sevření a puštění.", 10:"Rameno jelo, žádné sevření. Nespočteno kvůli posunuté základně – log to ani nepotvrdí, ani nevyvrátí.",
 11:"⚠ ROZPOR: úsudek ANO, ale poznámka i log říkají neúspěch (žádné trvalé sevření, 40 s do ručního ukončení). Pravděpodobně NE.",
 12:"Sedí: bez sevření.", 13:"Sedí: bez sevření.", 14:"Sedí: sevření a puštění.", 15:"Sedí: bez sevření.", 16:"Sedí: bez sevření.",
 17:"Sedí: sevření a puštění.", 18:"Sedí: sevření a puštění.", 19:"Sedí: bez sevření.",
 20:"Sedí: sevření 6,3 s, puštění 10,0 s; zbytek do 49,5 s stál nad miskou (jak píšeš).", 21:"Sedí: bez sevření (uchytil špatně).", 22:"Sedí: sevření a puštění.",
}
r += 2
wk.cell(r, 1, "3) Baseline – řádek po řádku (indicie z telemetrie zátěže gripperu, NE důkaz úspěchu)").font = S_FONT
r += 1
header(r, ["Řádek", "Soubor", "Doba [s]", "Váš úsudek", "Pohyb ramene [°]", "Max. zátěž", "Sevřel v [s]", "Držel [s]", "Puštění v [s]", "Indicie", "Shoda s úsudkem", "Komentář"] + ["", ""])
wk.row_dimensions[r].height = 42
for row in range(5, 23):
    b = base_info[row]; user = wbz.cell(row, 8).value
    proxy = "sevřel + pustil" if b["release"] else "bez trvalého sevření"
    agree = (user == "ANO") == bool(b["release"])
    r += 1
    put(r, [row, b["file"], b["dur"], user, round(b["move"], 1), round(b["maxload"]), round(b["grasp"], 1) if b["grasp"] else "-", round(b["held"], 1),
            round(b["release"], 1) if b["release"] else "-", proxy, "shoda" if agree else "NESEDÍ (viz komentář)", BCOMM[row]],
        fill=WARN if row == 11 else None)
    wk.merge_cells(start_row=r, start_column=12, end_row=r, end_column=14)
    wk.row_dimensions[r].height = 32

widths = [9, 24, 16, 34, 9, 9, 16, 12, 11, 10, 12, 26, 14, 60]
for i, w in enumerate(widths, 1):
    wk.column_dimensions[get_column_letter(i)].width = w
wk.freeze_panes = "A4"

wb.save(DST)
print("uloženo:", DST)
json.dump({"summ": summ, "orch": {k: {kk: vv for kk, vv in v.items() if kk != "r"} for k, v in orch_info.items()},
           "base": base_info}, open(SP + "fill_result.json", "w", encoding="utf-8"), ensure_ascii=False, default=str)
for k, (n, kk, avg) in summ.items():
    lo, hi = wilson(kk, n)
    print(f"{k}: N={n} úspěchů={kk} = {kk/n*100:.1f} %  Wilson [{lo*100:.0f}–{hi*100:.0f}] %  prům. doba {avg:.0f} s")
