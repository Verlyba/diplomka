"""60ep běhy (řádky 26–42): doplnění automatických sloupců, list Kontrola a porovnání s 120ep orchestrací (cs15)."""
import json, os, sys, io, math
import numpy as np, openpyxl
from PIL import Image, ImageDraw, ImageFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
ROOT = "C:/Users/green/diplomka/"
SP = "C:/Users/green/AppData/Local/Temp/claude/C--Users-green-diplomka/3092d090-a58d-4566-873f-9ee439d802b4/scratchpad/"
SRC = r"D:\_WinZakladniSlozky\zaznam_behu_2_doplneno.xlsx"
DST = r"D:\_WinZakladniSlozky\zaznam_behu_2_doplneno_60ep.xlsx"
OUT_PNG = ROOT + "poznamky/mapa_60ep_vs_120ep.png"

MAP = {26: "20260926-212027", 27: "20260926-212145", 28: "20260926-212754", 29: "20260926-213634", 30: "20260926-213929",
       31: "20260926-214823", 32: "20260926-215721", 33: "20260926-220442", 34: "20260926-220934",
       36: "20260926-221749", 37: "20260926-223157", 38: "20260926-223329", 39: "20260926-224235", 40: "20260926-224404",
       41: "20260926-224808", 42: "20260926-225113"}
NOLOG_ROW = 35
CRASH_ROW, CRASH_ID = 36, "20260926-221749"
REAL_TO_HOMING, HOMING_MEAN = 46.6, 8.3
ADJ = round(REAL_TO_HOMING + HOMING_MEAN, 1)
FINAL = {27: "v misce", 28: "mimo misku", 29: "v misce", 30: "mimo misku", 31: "mimo misku", 32: "mimo misku (v záběru ruka)", 33: "mimo misku",
         34: "mimo misku", 36: "v misce", 37: "v misce", 38: "v misce", 39: "v misce", 40: "mimo misku (u horního okraje desky)", 41: "v misce",
         42: "mimo misku (mimo desku vlevo nahoře)"}

wb = openpyxl.load_workbook(SRC)
wo = wb["Orchestrace"]
pos60 = json.load(open(SP + "pos60.json"))
info = {}
for row, rid in MAP.items():
    r = json.load(open(ROOT + f"runs/{rid}.json", encoding="utf-8"))
    cs = r.get("cost_summary") or {}
    st = r["steps"]
    assert "60ep" in json.dumps(r["catalog"]), rid
    cond = "Orchestrace (bez re-planu)" if r["config"].get("max_replans") == 0 else "Orchestrace"
    d, t = f"{rid[:4]}-{rid[4:6]}-{rid[6:8]}", f"{rid[9:11]}:{rid[11:13]}"
    dur = ADJ if row == CRASH_ROW else round(float(r["duration_s"]), 1)
    vals = {1: rid, 2: d, 3: t, 4: cond, 6: "60", 7: dur, 8: max([s.get("replan", 0) for s in st] or [0]),
            9: cs.get("planner_calls", 0), 10: cs.get("inspector_calls", 0), 11: "ANO" if r.get("success") else "NE"}
    for col, v in vals.items():
        assert wo.cell(row, col).value in (None, ""), (row, col, wo.cell(row, col).value)
        wo.cell(row, col, v)
    if wo.cell(row, 5).value in (None, ""): wo.cell(row, 5, "60 cs15")
    seq = " ".join(f"{ {'catch_cube':'C','carry_cube':'K','homing':'H'}.get(s['step'], '?')}{'+' if s['success'] else '-'}" for s in st) or "(žádný krok)"
    info[row] = dict(id=rid, seq=seq, sys=vals[11], dur=vals[7], raw=r["duration_s"], rep=vals[8], ceo=vals[9], vlm=vals[10],
                     n_catch=sum(1 for s in st if s["step"] == "catch_cube"), pos=pos60.get(rid))
# řádek 35: bez záznamu – jen podmínka, model a epizody, aby se úsudek uživatele započítal do Souhrnu
for col, v in {4: "Orchestrace", 6: "60"}.items():
    assert wo.cell(NOLOG_ROW, col).value in (None, ""); wo.cell(NOLOG_ROW, col, v)
if wo.cell(NOLOG_ROW, 5).value in (None, ""): wo.cell(NOLOG_ROW, 5, "60 cs15")

# řádek 36: pád démona – upravená doba, výslovně označená
note = wo.cell(CRASH_ROW, 14).value or ""
disc = (f" [ÚPRAVA DAT (doplnil asistent na pokyn uživatele): běh {CRASH_ID} skončil chybou démona před homingem (výměna modelu se nepotvrdila, po 180 s restart a pád "
        f"UnboundLocalError; v záznamu error „Daemon neběží.“) a homing neproběhl. Doba behu {ADJ} s = skutečných {REAL_TO_HOMING} s do konce ověření carry_cube + "
        f"{HOMING_MEAN} s průměrná doba homingu (výměna modelu + krok + ověření, 16 dokončených homingů). Uložená celková doba záznamu je {info[CRASH_ROW]['raw']} s "
        f"(zahrnuje ~3 min čekání na chybu). Počty volání CEO/VLM jsou naměřené, bez homingu. Verdikt systému NE odpovídá záznamu.]")
wo.cell(CRASH_ROW, 14, (note + disc).strip()); wo.cell(CRASH_ROW, 14).alignment = Alignment(wrap_text=True, vertical="top")

# ── 60ep podle poloh + 120ep podle poloh
rows120 = json.load(open(SP + "rows.json"))                       # 16 poloh, orchestrace 120ep cs15
ROW_OF_POS = {}                                                   # poloha -> řádek 60ep
for row, i in info.items():
    if i["pos"] and row != 26: ROW_OF_POS[i["pos"][1]] = row
ROW_OF_POS[5] = NOLOG_ROW                                         # poloha 5 nemá běh (vyloučením)
o60 = {}
for k in range(1, 17):
    row = ROW_OF_POS[k]; user = wo.cell(row, 13).value
    i = info.get(row)
    o60[k] = dict(row=row, user=user, n_catch=(i["n_catch"] if i else None), rep=(i["rep"] if i else None), dur=(i["dur"] if i else None),
                  id=(i["id"] if i else ""), sys=(i["sys"] if i else ""))
A = [rows120[k - 1]["o_user"] == "ANO" for k in range(1, 17)]; B = [o60[k]["user"] == "ANO" for k in range(1, 17)]
both = sum(a and b for a, b in zip(A, B)); only120 = sum(a and not b for a, b in zip(A, B)); only60 = sum(b and not a for a, b in zip(A, B)); none = sum(not a and not b for a, b in zip(A, B))
def mcnemar(b, c):
    n = b + c; k = min(b, c); return min(1.0, 2 * sum(math.comb(n, i) for i in range(0, k + 1)) / 2 ** n) if n else 1.0
def fisher(a, n1, b, n2):
    K = a + b; N = n1 + n2; p_obs = math.comb(n1, a) * math.comb(n2, b) / math.comb(N, K); tot = 0
    for x in range(max(0, K - n2), min(n1, K) + 1):
        p = math.comb(n1, x) * math.comb(n2, K - x) / math.comb(N, K)
        if p <= p_obs + 1e-12: tot += p
    return tot
def wilson(k, n, z=1.959963985):
    p = k / n; c = (p + z * z / (2 * n)) / (1 + z * z / n); m = (z / (1 + z * z / n)) * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)); return round(max(0, c - m) * 100), round(min(1, c + m) * 100)
s120, s60 = sum(A), sum(B)
first120 = sum(1 for r in rows120 if r["o_user"] == "ANO" and r["o_rep"] == 0)
first60 = sum(1 for k in range(1, 17) if o60[k]["user"] == "ANO" and (o60[k]["rep"] == 0 or k == 5))
print(f"120ep {s120}/16 {wilson(s120,16)} | 60ep {s60}/16 {wilson(s60,16)} | obě {both}, jen 120ep {only120}, jen 60ep {only60}, ani jedna {none} | McNemar p={mcnemar(only120, only60):.2f} | Fisher p={fisher(s120,16,s60,16):.2f}")
print("úspěch bez re-plánu: 120ep", first120, "/16 | 60ep", first60, "/16 (řádek 35 podle poznámky 'na první pokus')")

# ── list Kontrola
wk = wb["Kontrola"]
FONT = "Arial"
Hf = PatternFill("solid", fgColor="1F4E78"); Hn = Font(name=FONT, size=10, bold=True, color="FFFFFF"); Nn = Font(name=FONT, size=10); Bn = Font(name=FONT, size=10, bold=True)
Sn = Font(name=FONT, size=12, bold=True, color="1F4E78")
thin = Side(style="thin", color="D9D9D9"); BR = Border(left=thin, right=thin, top=thin, bottom=thin)
warn, okf = PatternFill("solid", fgColor="FCE4D6"), PatternFill("solid", fgColor="E2EFDA")
for r in range(1, wk.max_row + 1):
    if str(wk.cell(r, 1).value or "").startswith("⚠ Baseline řádek 11"):
        wk.cell(r, 1, "Baseline řádek 11 (vyřešeno)")
        wk.cell(r, 2, "Úsudek původně ANO odporoval poznámce i telemetrii; uživatel ho opravil na NE (uloženo). Se změnou vychází baseline 7/16.")
        for c in range(1, 15): wk.cell(r, c).fill = PatternFill(fill_type=None)

def block_title(r, text): wk.cell(r, 1, text).font = Sn
def put_find(r, t, txt, w=warn):
    wk.cell(r, 1, t).font = Bn; wk.cell(r, 2, txt).font = Nn
    for c in range(1, 15): wk.cell(r, c).border = BR
    wk.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top"); wk.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
    wk.merge_cells(start_row=r, start_column=2, end_row=r, end_column=14)
    wk.row_dimensions[r].height = max(30, 15 * math.ceil(len(txt) / 150))
    if w and t.startswith("⚠"):
        for c in range(1, 15): wk.cell(r, c).fill = w
r = wk.max_row + 2
block_title(r, "4) Běhy s 60 epizodami (řádky 26–42)")
r += 1
for i, h in enumerate(["Téma", "Nález"], 1):
    c = wk.cell(r, i, h); c.font = Hn; c.fill = Hf; c.border = BR
wk.merge_cells(start_row=r, start_column=2, end_row=r, end_column=14)
F4 = [
 ("Mapování", "Řádky 26–34 a 36–42 = běhy 20260926-212027 … 225113 v pořadí (řádek 26 je běh zastavený před prvním krokem; řádek 36 je běh 221749 s pádem démona). Potvrzeno vzorem úsudků, poznámkami a závěrečnými snímky (u všech 15 běhů se snímky sedí Váš úsudek). "
              "Polohy kostky: běhy pokrývají 15 z 16 křížků, vyloučením chybí č. 5 (195, 307)."),
 ("⚠ Řádek 35 bez záznamu", "Žádný běh v runs/ ani telemetry/ mu neodpovídá. Vyloučením patří poloze č. 5. Doplněno jen Podmínka, Model, Epizod = 60 (aby se Váš úsudek ANO započítal), ostatní sloupce zůstávají prázdné. Do časových průměrů nejde."),
 ("⚠ Řádek 36 – upravená doba", f"Doba behu {ADJ} s není naměřená: skutečných {REAL_TO_HOMING} s do konce carry_cube + {HOMING_MEAN} s průměrný homing (pád démona před homingem, homing neproběhl). Totéž je v Poznámce řádku 36 a v poznamky/DENIK.md. "
              "Naměřená doba záznamu je 262,6 s. Chcete-li čistě naměřená data, přepište G36 na 46,6 nebo řádek z časů vylučte."),
 ("Úsudek systému", "Rozdíl proti Vašemu úsudku u řádků 34 (běh 220934: systém ANO, kostka mimo misku), 36 (chyba démona, systém NE, kostka v misce) a 40 (běh 224404 jste zastavil, systém NE, sedí)."),
]
for t, txt in F4:
    r += 1; put_find(r, t, txt)
r += 2
hdr = ["Řádek", "ID běhu", "Poloha č.", "Kroky (C=catch K=carry H=homing)", "Doba v sešitě [s]", "Uložená doba [s]", "Re-plány", "Systém", "Váš úsudek", "Závěrečný snímek (moje čtení)", "Sedí úsudek?"]
for i, h in enumerate(hdr, 1):
    c = wk.cell(r, i, h); c.font = Hn; c.fill = Hf; c.border = BR; c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
wk.row_dimensions[r].height = 42
for row in sorted(list(MAP) + [NOLOG_ROW]):
    user = wo.cell(row, 13).value
    if row == NOLOG_ROW:
        vals = [row, "(bez záznamu)", "5 (vyloučením)", "-", "-", "-", "-", "-", user, "-", "nelze ověřit"]
        fill = warn
    else:
        i = info[row]; fin = FINAL.get(row, "(žádný snímek – běh zastaven před prvním krokem)")
        ok = True if row == 26 else (("v misce" in fin) == (user == "ANO"))
        p = i["pos"]; ptxt = f"{p[1]} ({p[2]:.0f} px)" if p else "-"
        vals = [row, i["id"], ptxt, i["seq"], i["dur"], i["raw"], i["rep"], i["sys"], user, fin, "ano" if ok else "NE – ZKONTROLUJ"]
        fill = warn if row == CRASH_ROW else None
    r += 1
    for j, v in enumerate(vals, 1):
        c = wk.cell(r, j, v); c.font = Nn; c.border = BR; c.alignment = Alignment(wrap_text=True, vertical="top")
        if fill: c.fill = fill

# ── 5) porovnání 60ep × 120ep
r += 3
block_title(r, "5) Porovnání: orchestrace cs15, 60 epizod × 120 epizod (16 poloh, úsudek uživatele)")
r += 1
for i, h in enumerate(["Téma", "Nález"], 1):
    c = wk.cell(r, i, h); c.font = Hn; c.fill = Hf; c.border = BR
wk.merge_cells(start_row=r, start_column=2, end_row=r, end_column=14)
F5 = [
 ("Úspěšnost", f"120 ep: {s120}/16 ({wilson(s120,16)[0]}–{wilson(s120,16)[1]} %) | 60 ep: {s60}/16 ({wilson(s60,16)[0]}–{wilson(s60,16)[1]} %). Fisherův test p = {fisher(s120,16,s60,16):.2f}, tedy žádný významný rozdíl."),
 ("Po polohách (páry)", f"Obě trefily {both}, jen 120 ep {only120}, jen 60 ep {only60}, ani jedna {none}. McNemar p = {mcnemar(only120, only60):.2f}. Poloha 5 u 60ep je jen Váš úsudek bez záznamu."),
 ("Bez re-plánu", f"Úspěch bez jediného re-plánu: 120 ep {first120}/16, 60 ep {first60}/16 (v tom řádek 35 podle poznámky „na první pokus“)."),
]
for t, txt in F5:
    r += 1; put_find(r, t, txt)
r += 2
hdr = ["Poloha č.", "x [px]", "y [px]", "120ep řádek", "120ep úsudek", "120ep úchopů", "120ep re-plánů", "60ep řádek", "60ep úsudek", "60ep úchopů", "60ep re-plánů", "Výsledek"]
for i, h in enumerate(hdr, 1):
    c = wk.cell(r, i, h); c.font = Hn; c.fill = Hf; c.border = BR; c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
wk.row_dimensions[r].height = 36
for k in range(1, 17):
    a = rows120[k - 1]; b = o60[k]
    res = "obě trefily" if (a["o_user"] == "ANO" and b["user"] == "ANO") else ("jen 120 ep" if a["o_user"] == "ANO" else ("jen 60 ep" if b["user"] == "ANO" else "ani jedna"))
    vals = [k, round(a["x"]), round(a["y"]), a["o_row"], a["o_user"], a["o_att"], a["o_rep"], b["row"], b["user"], b["n_catch"] if b["n_catch"] is not None else "-", b["rep"] if b["rep"] is not None else "-", res]
    r += 1
    for j, v in enumerate(vals, 1):
        c = wk.cell(r, j, v); c.font = Nn; c.border = BR; c.alignment = Alignment(horizontal="center")
    wk.cell(r, 5).fill = okf if a["o_user"] == "ANO" else warn; wk.cell(r, 9).fill = okf if b["user"] == "ANO" else warn

wb.save(DST)
print("uloženo:", DST)

# ── Souhrn 60ep podle sešitu
n = k = 0; durs = []
for rr in range(5, 60):
    if wo.cell(rr, 6).value == "60" and wo.cell(rr, 4).value == "Orchestrace" and wo.cell(rr, 12).value == "ANO":
        n += 1; k += wo.cell(rr, 13).value == "ANO"
        if isinstance(wo.cell(rr, 7).value, (int, float)): durs.append(wo.cell(rr, 7).value)
print(f"Souhrn Orchestrace 60ep: {k}/{n} = {k/n*100:.0f} %, Wilson {wilson(k, n)} | průměrná doba {sum(durs)/len(durs):.0f} s (N={len(durs)})")
dur120 = [r["o_dur"] for r in rows120]
print(f"průměrná doba 120ep cs15: {sum(dur120)/len(dur120):.0f} s (N=16)")
succ_dur60 = [wo.cell(rr, 7).value for rr in range(27, 43) if wo.cell(rr, 13).value == "ANO" and isinstance(wo.cell(rr, 7).value, (int, float))]
succ_dur120 = [r["o_dur"] for r in rows120 if r["o_user"] == "ANO"]
print(f"průměrná doba jen úspěšných: 60ep {sum(succ_dur60)/len(succ_dur60):.0f} s (N={len(succ_dur60)}), 120ep {sum(succ_dur120)/len(succ_dur120):.0f} s (N={len(succ_dur120)})")

# ── mapa: levá půlka 120ep, pravá 60ep
S = 2
cpos = json.load(open(SP + "cube_pos.json")); base = sorted(cpos)[:16]
imgs = [np.array(Image.open(ROOT + json.load(open(ROOT + f"runs/{r_}.json", encoding="utf-8"))["initial_images"][0]).convert("RGB")) for r_ in base]
bg = np.median(np.stack(imgs), axis=0).astype(np.uint8)
bgim = Image.blend(Image.fromarray(bg).resize((640 * S, 480 * S), Image.LANCZOS), Image.new("RGB", (640 * S, 480 * S), "white"), 0.45)
W, H = bgim.size; PAD = 200
canvas = Image.new("RGB", (W, H + PAD), "white"); canvas.paste(bgim, (0, 0)); d = ImageDraw.Draw(canvas, "RGBA")
def font(sz, bold=False):
    for f in (["C:/Windows/Fonts/arialbd.ttf"] if bold else []) + ["C:/Windows/Fonts/arial.ttf"]:
        try: return ImageFont.truetype(f, sz)
        except Exception: pass
    return ImageFont.load_default()
F, FB, FS = font(22), font(24, True), font(18)
G_, R_ = (46, 158, 79, 255), (214, 69, 69, 255)
demos = np.array(json.load(open(SP + "demos_pose_pos.json"))["P"])
for x, y in demos: d.ellipse([x * S - 3, y * S - 3, x * S + 3, y * S + 3], fill=(90, 90, 90, 110))
R = 25
for k in range(1, 17):
    a = rows120[k - 1]; b = o60[k]; cx, cy = a["x"] * S, a["y"] * S
    ca = G_ if a["o_user"] == "ANO" else R_; cb = G_ if b["user"] == "ANO" else R_
    d.pieslice([cx - R, cy - R, cx + R, cy + R], 90, 270, fill=ca); d.pieslice([cx - R, cy - R, cx + R, cy + R], 270, 90, fill=cb)
    d.ellipse([cx - R, cy - R, cx + R, cy + R], outline=(255, 255, 255, 255), width=3)
    if k == 5: d.ellipse([cx - R - 8, cy - R - 8, cx + R + 8, cy + R + 8], outline=(60, 60, 60, 255), width=3)
    lab = str(k); tw = d.textlength(lab, font=FB)
    d.rectangle([cx - tw / 2 - 4, cy - R - 30, cx + tw / 2 + 4, cy - R - 4], fill=(255, 255, 255, 230)); d.text((cx - tw / 2, cy - R - 31), lab, fill=(20, 20, 20, 255), font=FB)
y0 = H + 12
d.text((16, y0), "Orchestrace cs15: 120 epizod × 60 epizod (poloha kostky na začátku, úsudek uživatele)", fill=(20, 20, 20, 255), font=FB)
d.pieslice([16, y0 + 44, 52, y0 + 80], 90, 270, fill=G_); d.pieslice([16, y0 + 44, 52, y0 + 80], 270, 90, fill=R_); d.ellipse([16, y0 + 44, 52, y0 + 80], outline=(120, 120, 120, 255), width=2)
d.text((66, y0 + 48), "levá půlka = 120 epizod, pravá půlka = 60 epizod; zelená = trefil, červená = netrefil", fill=(20, 20, 20, 255), font=F)
d.ellipse([16, y0 + 96, 40, y0 + 120], outline=(60, 60, 60, 255), width=3)
d.text((66, y0 + 96), "kroužek = poloha č. 5: u 60 epizod jen váš úsudek, bez záznamu běhu", fill=(20, 20, 20, 255), font=F)
d.text((16, y0 + 140), f"120 ep: {s120}/16 | 60 ep: {s60}/16 | obě trefily {both}, jen 120 ep {only120}, jen 60 ep {only60}, ani jedna {none}", fill=(20, 20, 20, 255), font=F)
canvas.save(OUT_PNG); print("obrázek:", OUT_PNG)
json.dump({str(k): o60[k] for k in o60}, open(SP + "o60.json", "w"), ensure_ascii=False)
