"""Sestaví balík dat pro psací modely: mereni/ (+ analyza/). Nic se nemaže, jen kopíruje a exportuje."""
import csv, glob, json, math, os, shutil, sys, io
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
ROOT = "C:/Users/green/diplomka/"
SP = "C:/Users/green/AppData/Local/Temp/claude/C--Users-green-diplomka/3092d090-a58d-4566-873f-9ee439d802b4/scratchpad/"
XLSX = r"D:\_WinZakladniSlozky\zaznam_behu_2_doplneno_60ep.xlsx"
M = ROOT + "mereni/"
FIRST = "20260926-164147"                          # první započítaný běh orchestrace

for d in ("", "zaznamy_behu", "telemetrie", "snimky", "demonstrace"):
    os.makedirs(M + d, exist_ok=True)
os.makedirs(ROOT + "analyza", exist_ok=True)

# ── 1) sešit a obrázky
shutil.copy2(XLSX, M + "zaznam_behu_2.xlsx")
shutil.copy2(ROOT + "poznamky/mapa_polohy_kostky.png", M + "mapa_120ep_orchestrace_vs_baseline.png")
shutil.copy2(ROOT + "poznamky/mapa_60ep_vs_120ep.png", M + "mapa_60ep_vs_120ep.png")
shutil.copy2(ROOT + "poznamky/poloha_kostky_start.xlsx", M + "poloha_kostky_start_120ep_baseline.xlsx")

# ── 2) záznamy běhů, telemetrie, snímky (jen měřené běhy od 164147; adresáře se jmenují jinak než v .gitignore)
n_runs = n_tel = n_img = 0
for f in sorted(glob.glob(ROOT + "runs/20260926-*.json")):
    rid = os.path.basename(f)[:-5]
    if rid >= FIRST:
        shutil.copy2(f, M + "zaznamy_behu/" + os.path.basename(f)); n_runs += 1
        d = ROOT + f"images/{rid}"
        if os.path.isdir(d):
            shutil.copytree(d, M + f"snimky/{rid}", dirs_exist_ok=True); n_img += len(os.listdir(d))
for f in sorted(glob.glob(ROOT + "telemetry/20260926-*.jsonl")):
    if os.path.basename(f)[:15] >= "20260926-164000":
        shutil.copy2(f, M + "telemetrie/" + os.path.basename(f)); n_tel += 1
print(f"záznamů běhů {n_runs}, telemetrií {n_tel}, snímků {n_img}")

# ── 3) CSV z listů sešitu
wb = openpyxl.load_workbook(XLSX)
def sheet_csv(name, out, first_col, ncols):
    ws = wb[name]
    hdr = [ws.cell(4, c).value for c in range(1, ncols + 1)]
    with open(out, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh); w.writerow(hdr)
        for r in range(5, ws.max_row + 1):
            vals = [ws.cell(r, c).value for c in range(1, ncols + 1)]
            if any(v not in (None, "") for v in vals[first_col:]) or vals[0]:
                w.writerow([r] + ["" if v is None else v for v in vals]) if False else w.writerow(["" if v is None else v for v in vals] + [r])
    return hdr
h1 = sheet_csv("Orchestrace", M + "orchestrace_behy.csv", 0, 14)
h2 = sheet_csv("Baseline", M + "baseline_pokusy.csv", 0, 9)
# poslední sloupec „radek_v_sesitu“ přidán jako hlavička
for fn, hdr in (("orchestrace_behy.csv", h1), ("baseline_pokusy.csv", h2)):
    p = M + fn; rows = list(csv.reader(open(p, encoding="utf-8")))
    rows[0] = rows[0] + ["radek_v_sesitu"]
    with open(p, "w", newline="", encoding="utf-8") as fh: csv.writer(fh).writerows(rows)

# ── 4) shrnutí (Souhrn v sešitu jsou vzorce bez uložených hodnot, proto zvlášť)
def wilson(k, n, z=1.959963985):
    p = k / n; c = (p + z * z / (2 * n)) / (1 + z * z / n); m = (z / (1 + z * z / n)) * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)); return max(0, c - m), min(1, c + m)
wo, wbz = wb["Orchestrace"], wb["Baseline"]
summ = []
def add(label, ep, rows_iter):
    n = k = 0; d = []
    for ok, dur in rows_iter:
        n += 1; k += ok
        if isinstance(dur, (int, float)): d.append(dur)
    lo, hi = wilson(k, n)
    summ.append([label, ep, n, k, round(k / n, 3), round(lo, 3), round(hi, 3), round(sum(d) / len(d), 1) if d else "", len(d)])
add("orchestrace cs15", 120, ((wo.cell(r, 13).value == "ANO", wo.cell(r, 7).value) for r in range(5, 21) if wo.cell(r, 6).value == "120" and wo.cell(r, 12).value == "ANO" and wo.cell(r, 5).value != "120 cs 100"))
add("orchestrace cs15", 60, ((wo.cell(r, 13).value == "ANO", wo.cell(r, 7).value) for r in range(26, 43) if wo.cell(r, 6).value == "60" and wo.cell(r, 12).value == "ANO"))
add("baseline (cs15)", 120, ((wbz.cell(r, 8).value == "ANO", wbz.cell(r, 6).value) for r in range(5, 23) if wbz.cell(r, 7).value == "ANO"))
with open(M + "souhrn.csv", "w", newline="", encoding="utf-8") as fh:
    w = csv.writer(fh); w.writerow(["podminka", "epizod", "n_pocita_se", "uspechu_usudek_uzivatele", "uspesnost", "wilson_dolni", "wilson_horni", "prumerna_doba_s", "n_s_dobou"]); w.writerows(summ)
print("souhrn:", summ)

# ── 5) polohy × výsledky (16 poloh)
rows120 = json.load(open(SP + "rows.json")); o60 = json.load(open(SP + "o60.json"))
demos = json.load(open(SP + "demos_pose_pos.json")); P = demos["P"]; Q = demos["Q"]
import numpy as np
Pn = np.array(P)
def cov(x, y, n): return int((np.hypot(Pn[:n, 0] - x, Pn[:n, 1] - y) < 60).sum())
base_est = json.load(open(SP + "base_pos.json"))
with open(M + "polohy_a_vysledky.csv", "w", newline="", encoding="utf-8") as fh:
    w = csv.writer(fh)
    w.writerow(["poloha", "x_px", "y_px", "demonstraci_do_60px_z_120ep", "demonstraci_do_60px_z_prvnich_60ep",
                "orch120_radek", "orch120_run", "orch120_usudek", "orch120_uchopu", "orch120_replanu", "orch120_doba_s",
                "orch60_radek", "orch60_run", "orch60_usudek", "orch60_uchopu", "orch60_replanu", "orch60_doba_s", "orch60_poznamka",
                "baseline_radek", "baseline_soubor", "baseline_usudek", "baseline_doba_s", "baseline_prirazeni", "baseline_odhad_x_px", "baseline_odhad_y_px"])
    for k in range(1, 17):
        a = rows120[k - 1]; b = o60[str(k)]; br = a["b_row"]; est = a.get("b_est")
        nota = ""
        if b["row"] == 35: nota = "bez záznamu běhu (jen úsudek uživatele)"
        if b["row"] == 36: nota = "doba upravena: 46.6 s skutečných + 8.3 s průměrný homing; běh skončil pádem démona před homingem"
        w.writerow([k, round(a["x"]), round(a["y"]), cov(a["x"], a["y"], 120), cov(a["x"], a["y"], 60),
                    a["o_row"], a["o_id"], a["o_user"], a["o_att"], a["o_rep"], a["o_dur"],
                    b["row"], b["id"], b["user"], "" if b["n_catch"] is None else b["n_catch"], "" if b["rep"] is None else b["rep"], "" if b["dur"] is None else b["dur"], nota,
                    br, a["b_file"], a["b_user"], wbz.cell(br, 6).value, "nejisté (blok poloh 5-9)" if a["b_uncertain"] else "jisté",
                    "" if not est else round(est[0]), "" if not est else round(est[1])])

# ── 6) demonstrace: poloha kostky a koncová póza ramene každé epizody
with open(M + "demonstrace/pozice_demonstraci.csv", "w", newline="", encoding="utf-8") as fh:
    w = csv.writer(fh); w.writerow(["epizoda", "kostka_x_px", "kostka_y_px", "pan_deg", "shoulder_lift_deg", "elbow_flex_deg", "wrist_flex_deg", "wrist_roll_deg"])
    for i, (p, q) in enumerate(zip(P, Q)): w.writerow([i, round(p[0], 1), round(p[1], 1)] + [round(x, 1) for x in q])

# ── 7) skripty
SCRIPTS = ["pose2pos.py", "base_pos.py", "base_check.py", "coverage.py", "cube_pos.py", "export_map.py", "fill_xlsx.py", "fill_60ep_v2.py", "replay_release.py",
           "carry_end.py", "carry_traces.py", "orch_summary.py", "list_new.py", "list_base.py", "montage.py", "finals.py", "build_mereni.py"]
for s in SCRIPTS:
    src = SP + s
    if os.path.exists(src): shutil.copy2(src, ROOT + "analyza/" + s)
print("skriptů:", len(os.listdir(ROOT + "analyza")))
